#!/usr/bin/env python3
"""
BUKU KAS KELOMPOK
Aplikasi pencatatan kas untuk kelompok masyarakat (PKK/Karang Taruna/Kelompok Tani/PPG/dll)
Single file Python application dengan SQLite database
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import sqlite3
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
import os
import locale

# Set locale untuk format angka Indonesia
try:
    locale.setlocale(locale.LC_ALL, 'id_ID.UTF-8')
except:
    try:
        locale.setlocale(locale.LC_ALL, '')
    except:
        pass

# Konstanta
DB_FILE = "kas_kelompok.db"
WINDOW_WIDTH = 1050
WINDOW_HEIGHT = 700

# Jenis transaksi
JENIS_TRANSAKSI = {
    'TRM': 'Terima Dana',
    'BLJ': 'Belanja Tunai',
    'STR': 'Setor ke Bank',
    'TRK': 'Tarik dari Bank',
    'TRF': 'Transfer/Bayar via Bank',
    'ADM': 'Biaya Admin Bank',
    'BNG': 'Bunga/Jasa Giro',
    'SDS': 'Setor ke Desa',
    'SDR': 'Setor ke Daerah'
}

# Kategori warna transaksi
WARNA_TRANSAKSI = {
    'masuk': '#d4edda',      # Hijau muda (TRM, BNG)
    'keluar': '#f8d7da',     # Merah muda (BLJ, TRF, ADM)
    'pindah': '#cce5ff',     # Biru muda (STR, TRK)
    'setor': '#fff3cd'       # Kuning (SDS, SDR)
}


def format_rupiah(nilai):
    """Format angka menjadi format Rupiah"""
    if nilai is None:
        return "Rp 0"
    try:
        nilai = float(nilai)
        if nilai < 0:
            return f"-Rp {abs(nilai):,.0f}".replace(",", ".")
        return f"Rp {nilai:,.0f}".replace(",", ".")
    except:
        return "Rp 0"


def parse_rupiah(text):
    """Parse string rupiah menjadi float"""
    if not text:
        return 0.0
    text = str(text).replace("Rp", "").replace(".", "").replace(",", ".").strip()
    try:
        return float(text)
    except:
        return 0.0


class Database:
    """Kelas untuk mengelola koneksi dan operasi database"""

    def __init__(self, db_file):
        self.db_file = db_file
        self.conn = None
        self.init_database()

    def get_connection(self):
        """Mendapatkan koneksi database"""
        if self.conn is None:
            self.conn = sqlite3.connect(self.db_file)
            self.conn.row_factory = sqlite3.Row
        return self.conn

    def init_database(self):
        """Inisialisasi struktur database"""
        conn = self.get_connection()
        cursor = conn.cursor()

        # Tabel pengaturan
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS pengaturan (
                id INTEGER PRIMARY KEY,
                nama_kelompok TEXT DEFAULT 'Nama Kelompok',
                alamat TEXT DEFAULT '',
                nama_bendahara TEXT DEFAULT ''
            )
        ''')

        # Insert default pengaturan jika belum ada
        cursor.execute('SELECT COUNT(*) FROM pengaturan')
        if cursor.fetchone()[0] == 0:
            cursor.execute('''
                INSERT INTO pengaturan (nama_kelompok, alamat, nama_bendahara)
                VALUES ('Nama Kelompok', '', '')
            ''')

        # Tabel pos_anggaran
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS pos_anggaran (
                id INTEGER PRIMARY KEY,
                kode_pos TEXT UNIQUE NOT NULL,
                nama_pos TEXT NOT NULL,
                target_anggaran REAL DEFAULT 0,
                keterangan TEXT DEFAULT ''
            )
        ''')

        # Tabel buku_kas
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS buku_kas (
                id INTEGER PRIMARY KEY,
                tanggal TEXT NOT NULL,
                uraian TEXT NOT NULL,
                jenis_transaksi TEXT NOT NULL,
                jumlah REAL NOT NULL,
                sumber_kas TEXT,
                pos_id INTEGER,
                arah_pos TEXT,
                saldo_tunai_setelah REAL DEFAULT 0,
                saldo_bank_setelah REAL DEFAULT 0,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (pos_id) REFERENCES pos_anggaran(id)
            )
        ''')

        # Tabel rekonsiliasi
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS rekonsiliasi (
                id INTEGER PRIMARY KEY,
                tanggal TEXT NOT NULL,
                tunai_fisik REAL DEFAULT 0,
                tunai_sistem REAL DEFAULT 0,
                selisih_tunai REAL DEFAULT 0,
                bank_fisik REAL DEFAULT 0,
                bank_sistem REAL DEFAULT 0,
                selisih_bank REAL DEFAULT 0,
                keterangan TEXT DEFAULT '',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        ''')

        conn.commit()

    def execute(self, query, params=()):
        """Eksekusi query dengan parameter"""
        conn = self.get_connection()
        cursor = conn.cursor()
        cursor.execute(query, params)
        conn.commit()
        return cursor

    def fetchall(self, query, params=()):
        """Fetch semua hasil query"""
        conn = self.get_connection()
        cursor = conn.cursor()
        cursor.execute(query, params)
        return cursor.fetchall()

    def fetchone(self, query, params=()):
        """Fetch satu hasil query"""
        conn = self.get_connection()
        cursor = conn.cursor()
        cursor.execute(query, params)
        return cursor.fetchone()

    def get_pengaturan(self):
        """Mendapatkan pengaturan aplikasi"""
        return self.fetchone('SELECT * FROM pengaturan LIMIT 1')

    def update_pengaturan(self, nama_kelompok, alamat, nama_bendahara):
        """Update pengaturan aplikasi"""
        self.execute('''
            UPDATE pengaturan SET nama_kelompok=?, alamat=?, nama_bendahara=?
            WHERE id=1
        ''', (nama_kelompok, alamat, nama_bendahara))

    def get_saldo_tunai(self):
        """Mendapatkan saldo tunai terkini"""
        result = self.fetchone('''
            SELECT saldo_tunai_setelah FROM buku_kas
            ORDER BY tanggal DESC, id DESC LIMIT 1
        ''')
        return result['saldo_tunai_setelah'] if result else 0.0

    def get_saldo_bank(self):
        """Mendapatkan saldo bank terkini"""
        result = self.fetchone('''
            SELECT saldo_bank_setelah FROM buku_kas
            ORDER BY tanggal DESC, id DESC LIMIT 1
        ''')
        return result['saldo_bank_setelah'] if result else 0.0

    def get_all_pos(self):
        """Mendapatkan semua pos anggaran"""
        return self.fetchall('SELECT * FROM pos_anggaran ORDER BY kode_pos')

    def get_pos_by_id(self, pos_id):
        """Mendapatkan pos anggaran berdasarkan ID"""
        return self.fetchone('SELECT * FROM pos_anggaran WHERE id=?', (pos_id,))

    def add_pos(self, kode_pos, nama_pos, target_anggaran, keterangan):
        """Menambah pos anggaran baru"""
        self.execute('''
            INSERT INTO pos_anggaran (kode_pos, nama_pos, target_anggaran, keterangan)
            VALUES (?, ?, ?, ?)
        ''', (kode_pos, nama_pos, target_anggaran, keterangan))

    def update_pos(self, pos_id, kode_pos, nama_pos, target_anggaran, keterangan):
        """Update pos anggaran"""
        self.execute('''
            UPDATE pos_anggaran SET kode_pos=?, nama_pos=?, target_anggaran=?, keterangan=?
            WHERE id=?
        ''', (kode_pos, nama_pos, target_anggaran, keterangan, pos_id))

    def delete_pos(self, pos_id):
        """Hapus pos anggaran"""
        self.execute('DELETE FROM pos_anggaran WHERE id=?', (pos_id,))

    def pos_has_transactions(self, pos_id):
        """Cek apakah pos memiliki transaksi"""
        result = self.fetchone('SELECT COUNT(*) as cnt FROM buku_kas WHERE pos_id=?', (pos_id,))
        return result['cnt'] > 0

    def get_pos_summary(self, pos_id):
        """Mendapatkan ringkasan pos (masuk, keluar)"""
        masuk = self.fetchone('''
            SELECT COALESCE(SUM(jumlah), 0) as total FROM buku_kas
            WHERE pos_id=? AND arah_pos='masuk'
        ''', (pos_id,))
        keluar = self.fetchone('''
            SELECT COALESCE(SUM(jumlah), 0) as total FROM buku_kas
            WHERE pos_id=? AND arah_pos='keluar'
        ''', (pos_id,))
        return {
            'masuk': masuk['total'] if masuk else 0,
            'keluar': keluar['total'] if keluar else 0
        }

    def get_all_transaksi(self, tgl_dari=None, tgl_sampai=None):
        """Mendapatkan semua transaksi dengan filter tanggal"""
        query = '''
            SELECT bk.*, pa.kode_pos, pa.nama_pos
            FROM buku_kas bk
            LEFT JOIN pos_anggaran pa ON bk.pos_id = pa.id
        '''
        params = []

        if tgl_dari and tgl_sampai:
            query += ' WHERE bk.tanggal BETWEEN ? AND ?'
            params = [tgl_dari, tgl_sampai]
        elif tgl_dari:
            query += ' WHERE bk.tanggal >= ?'
            params = [tgl_dari]
        elif tgl_sampai:
            query += ' WHERE bk.tanggal <= ?'
            params = [tgl_sampai]

        query += ' ORDER BY bk.tanggal ASC, bk.id ASC'
        return self.fetchall(query, params)

    def get_transaksi_by_id(self, trans_id):
        """Mendapatkan transaksi berdasarkan ID"""
        return self.fetchone('SELECT * FROM buku_kas WHERE id=?', (trans_id,))

    def add_transaksi(self, tanggal, uraian, jenis, jumlah, sumber_kas, pos_id, arah_pos):
        """Menambah transaksi baru"""
        # Hitung saldo sebelumnya
        saldo_tunai = self.get_saldo_tunai()
        saldo_bank = self.get_saldo_bank()

        # Hitung perubahan saldo berdasarkan jenis transaksi
        if jenis == 'TRM':  # Terima Dana
            if sumber_kas == 'tunai':
                saldo_tunai += jumlah
            else:
                saldo_bank += jumlah
        elif jenis == 'BLJ':  # Belanja Tunai
            saldo_tunai -= jumlah
        elif jenis == 'STR':  # Setor ke Bank
            saldo_tunai -= jumlah
            saldo_bank += jumlah
        elif jenis == 'TRK':  # Tarik dari Bank
            saldo_bank -= jumlah
            saldo_tunai += jumlah
        elif jenis == 'TRF':  # Transfer via Bank
            saldo_bank -= jumlah
        elif jenis == 'ADM':  # Biaya Admin Bank
            saldo_bank -= jumlah
        elif jenis == 'BNG':  # Bunga/Jasa Giro
            saldo_bank += jumlah
        elif jenis in ('SDS', 'SDR'):  # Setor ke Desa/Daerah
            if sumber_kas == 'tunai':
                saldo_tunai -= jumlah
            else:
                saldo_bank -= jumlah

        cursor = self.execute('''
            INSERT INTO buku_kas (tanggal, uraian, jenis_transaksi, jumlah, sumber_kas,
                                  pos_id, arah_pos, saldo_tunai_setelah, saldo_bank_setelah, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (tanggal, uraian, jenis, jumlah, sumber_kas, pos_id, arah_pos,
              saldo_tunai, saldo_bank, datetime.now().isoformat()))

        # Recalculate saldo untuk transaksi setelahnya (jika insert di tengah)
        self.recalculate_saldo_after(tanggal, cursor.lastrowid)

        return cursor.lastrowid

    def update_transaksi(self, trans_id, tanggal, uraian, jenis, jumlah, sumber_kas, pos_id, arah_pos):
        """Update transaksi"""
        self.execute('''
            UPDATE buku_kas SET tanggal=?, uraian=?, jenis_transaksi=?, jumlah=?,
                               sumber_kas=?, pos_id=?, arah_pos=?
            WHERE id=?
        ''', (tanggal, uraian, jenis, jumlah, sumber_kas, pos_id, arah_pos, trans_id))

        # Recalculate semua saldo
        self.recalculate_all_saldo()

    def delete_transaksi(self, trans_id):
        """Hapus transaksi"""
        self.execute('DELETE FROM buku_kas WHERE id=?', (trans_id,))
        # Recalculate semua saldo
        self.recalculate_all_saldo()

    def recalculate_saldo_after(self, tanggal, after_id):
        """Recalculate saldo untuk transaksi setelah tanggal tertentu"""
        # Untuk kesederhanaan, recalculate semua
        self.recalculate_all_saldo()

    def recalculate_all_saldo(self):
        """Recalculate semua saldo berjalan"""
        transaksi = self.fetchall('SELECT * FROM buku_kas ORDER BY tanggal ASC, id ASC')

        saldo_tunai = 0.0
        saldo_bank = 0.0

        for t in transaksi:
            jenis = t['jenis_transaksi']
            jumlah = t['jumlah']
            sumber_kas = t['sumber_kas']

            if jenis == 'TRM':
                if sumber_kas == 'tunai':
                    saldo_tunai += jumlah
                else:
                    saldo_bank += jumlah
            elif jenis == 'BLJ':
                saldo_tunai -= jumlah
            elif jenis == 'STR':
                saldo_tunai -= jumlah
                saldo_bank += jumlah
            elif jenis == 'TRK':
                saldo_bank -= jumlah
                saldo_tunai += jumlah
            elif jenis == 'TRF':
                saldo_bank -= jumlah
            elif jenis == 'ADM':
                saldo_bank -= jumlah
            elif jenis == 'BNG':
                saldo_bank += jumlah
            elif jenis in ('SDS', 'SDR'):
                if sumber_kas == 'tunai':
                    saldo_tunai -= jumlah
                else:
                    saldo_bank -= jumlah

            self.execute('''
                UPDATE buku_kas SET saldo_tunai_setelah=?, saldo_bank_setelah=?
                WHERE id=?
            ''', (saldo_tunai, saldo_bank, t['id']))

    def validate_transaksi(self, jenis, jumlah, sumber_kas, exclude_id=None):
        """Validasi transaksi tidak menyebabkan saldo negatif"""
        saldo_tunai = self.get_saldo_tunai()
        saldo_bank = self.get_saldo_bank()

        # Jika edit, kembalikan nilai transaksi lama
        if exclude_id:
            old = self.get_transaksi_by_id(exclude_id)
            if old:
                old_jenis = old['jenis_transaksi']
                old_jumlah = old['jumlah']
                old_sumber = old['sumber_kas']

                # Reverse old transaction effect
                if old_jenis == 'TRM':
                    if old_sumber == 'tunai':
                        saldo_tunai -= old_jumlah
                    else:
                        saldo_bank -= old_jumlah
                elif old_jenis == 'BLJ':
                    saldo_tunai += old_jumlah
                elif old_jenis == 'STR':
                    saldo_tunai += old_jumlah
                    saldo_bank -= old_jumlah
                elif old_jenis == 'TRK':
                    saldo_bank += old_jumlah
                    saldo_tunai -= old_jumlah
                elif old_jenis == 'TRF':
                    saldo_bank += old_jumlah
                elif old_jenis == 'ADM':
                    saldo_bank += old_jumlah
                elif old_jenis == 'BNG':
                    saldo_bank -= old_jumlah
                elif old_jenis in ('SDS', 'SDR'):
                    if old_sumber == 'tunai':
                        saldo_tunai += old_jumlah
                    else:
                        saldo_bank += old_jumlah

        # Apply new transaction
        if jenis == 'BLJ':
            if saldo_tunai - jumlah < 0:
                return False, "Saldo tunai tidak mencukupi"
        elif jenis == 'STR':
            if saldo_tunai - jumlah < 0:
                return False, "Saldo tunai tidak mencukupi untuk setor ke bank"
        elif jenis == 'TRK':
            if saldo_bank - jumlah < 0:
                return False, "Saldo bank tidak mencukupi untuk tarik tunai"
        elif jenis == 'TRF':
            if saldo_bank - jumlah < 0:
                return False, "Saldo bank tidak mencukupi untuk transfer"
        elif jenis == 'ADM':
            if saldo_bank - jumlah < 0:
                return False, "Saldo bank tidak mencukupi untuk biaya admin"
        elif jenis in ('SDS', 'SDR'):
            if sumber_kas == 'tunai' and saldo_tunai - jumlah < 0:
                return False, "Saldo tunai tidak mencukupi untuk setoran"
            elif sumber_kas == 'bank' and saldo_bank - jumlah < 0:
                return False, "Saldo bank tidak mencukupi untuk setoran"

        return True, ""

    def get_all_rekonsiliasi(self):
        """Mendapatkan semua data rekonsiliasi"""
        return self.fetchall('SELECT * FROM rekonsiliasi ORDER BY tanggal DESC')

    def add_rekonsiliasi(self, tanggal, tunai_fisik, bank_fisik, keterangan):
        """Menambah data rekonsiliasi"""
        tunai_sistem = self.get_saldo_tunai()
        bank_sistem = self.get_saldo_bank()
        selisih_tunai = tunai_fisik - tunai_sistem
        selisih_bank = bank_fisik - bank_sistem

        self.execute('''
            INSERT INTO rekonsiliasi (tanggal, tunai_fisik, tunai_sistem, selisih_tunai,
                                      bank_fisik, bank_sistem, selisih_bank, keterangan, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (tanggal, tunai_fisik, tunai_sistem, selisih_tunai, bank_fisik, bank_sistem,
              selisih_bank, keterangan, datetime.now().isoformat()))

    def get_laporan_bulanan(self, bulan, tahun):
        """Mendapatkan data untuk laporan bulanan"""
        tgl_awal = f"{tahun}-{bulan:02d}-01"
        if bulan == 12:
            tgl_akhir = f"{tahun + 1}-01-01"
        else:
            tgl_akhir = f"{tahun}-{bulan + 1:02d}-01"

        # Saldo awal bulan
        saldo_awal = self.fetchone('''
            SELECT saldo_tunai_setelah, saldo_bank_setelah FROM buku_kas
            WHERE tanggal < ? ORDER BY tanggal DESC, id DESC LIMIT 1
        ''', (tgl_awal,))

        saldo_awal_tunai = saldo_awal['saldo_tunai_setelah'] if saldo_awal else 0
        saldo_awal_bank = saldo_awal['saldo_bank_setelah'] if saldo_awal else 0

        # Saldo akhir bulan
        saldo_akhir = self.fetchone('''
            SELECT saldo_tunai_setelah, saldo_bank_setelah FROM buku_kas
            WHERE tanggal >= ? AND tanggal < ? ORDER BY tanggal DESC, id DESC LIMIT 1
        ''', (tgl_awal, tgl_akhir))

        if not saldo_akhir:
            saldo_akhir_tunai = saldo_awal_tunai
            saldo_akhir_bank = saldo_awal_bank
        else:
            saldo_akhir_tunai = saldo_akhir['saldo_tunai_setelah']
            saldo_akhir_bank = saldo_akhir['saldo_bank_setelah']

        # Total penerimaan
        penerimaan = self.fetchone('''
            SELECT COALESCE(SUM(jumlah), 0) as total FROM buku_kas
            WHERE tanggal >= ? AND tanggal < ? AND jenis_transaksi IN ('TRM', 'BNG')
        ''', (tgl_awal, tgl_akhir))

        # Total pengeluaran
        pengeluaran = self.fetchone('''
            SELECT COALESCE(SUM(jumlah), 0) as total FROM buku_kas
            WHERE tanggal >= ? AND tanggal < ? AND jenis_transaksi IN ('BLJ', 'TRF', 'ADM')
        ''', (tgl_awal, tgl_akhir))

        # Total setoran desa/daerah
        setoran = self.fetchone('''
            SELECT COALESCE(SUM(jumlah), 0) as total FROM buku_kas
            WHERE tanggal >= ? AND tanggal < ? AND jenis_transaksi IN ('SDS', 'SDR')
        ''', (tgl_awal, tgl_akhir))

        # Rekap per pos
        rekap_pos = self.fetchall('''
            SELECT pa.id, pa.kode_pos, pa.nama_pos, pa.target_anggaran,
                   COALESCE(SUM(CASE WHEN bk.arah_pos = 'masuk' THEN bk.jumlah ELSE 0 END), 0) as masuk,
                   COALESCE(SUM(CASE WHEN bk.arah_pos = 'keluar' THEN bk.jumlah ELSE 0 END), 0) as keluar
            FROM pos_anggaran pa
            LEFT JOIN buku_kas bk ON pa.id = bk.pos_id
                AND bk.tanggal >= ? AND bk.tanggal < ?
            GROUP BY pa.id
            ORDER BY pa.kode_pos
        ''', (tgl_awal, tgl_akhir))

        # Daftar setoran
        daftar_setoran = self.fetchall('''
            SELECT bk.tanggal, bk.uraian, bk.jenis_transaksi, bk.jumlah,
                   pa.kode_pos, pa.nama_pos
            FROM buku_kas bk
            LEFT JOIN pos_anggaran pa ON bk.pos_id = pa.id
            WHERE bk.tanggal >= ? AND bk.tanggal < ? AND bk.jenis_transaksi IN ('SDS', 'SDR')
            ORDER BY bk.tanggal
        ''', (tgl_awal, tgl_akhir))

        return {
            'saldo_awal_tunai': saldo_awal_tunai,
            'saldo_awal_bank': saldo_awal_bank,
            'saldo_akhir_tunai': saldo_akhir_tunai,
            'saldo_akhir_bank': saldo_akhir_bank,
            'total_penerimaan': penerimaan['total'],
            'total_pengeluaran': pengeluaran['total'],
            'total_setoran': setoran['total'],
            'rekap_pos': rekap_pos,
            'daftar_setoran': daftar_setoran
        }


class TransaksiDialog(tk.Toplevel):
    """Dialog untuk tambah/edit transaksi"""

    def __init__(self, parent, db, transaksi=None):
        super().__init__(parent)
        self.db = db
        self.transaksi = transaksi
        self.result = None

        self.title("Edit Transaksi" if transaksi else "Tambah Transaksi")
        self.geometry("500x450")
        self.resizable(False, False)
        self.transient(parent)
        self.grab_set()

        self.create_widgets()

        if transaksi:
            self.load_data()

        self.center_window()

    def center_window(self):
        """Posisikan window di tengah"""
        self.update_idletasks()
        x = (self.winfo_screenwidth() - self.winfo_width()) // 2
        y = (self.winfo_screenheight() - self.winfo_height()) // 2
        self.geometry(f"+{x}+{y}")

    def create_widgets(self):
        """Membuat widget dialog"""
        main_frame = ttk.Frame(self, padding=20)
        main_frame.pack(fill=tk.BOTH, expand=True)

        # Tanggal
        ttk.Label(main_frame, text="Tanggal:").grid(row=0, column=0, sticky=tk.W, pady=5)
        self.tanggal_var = tk.StringVar(value=date.today().strftime("%Y-%m-%d"))
        self.tanggal_entry = ttk.Entry(main_frame, textvariable=self.tanggal_var, width=15)
        self.tanggal_entry.grid(row=0, column=1, sticky=tk.W, pady=5)
        ttk.Label(main_frame, text="(YYYY-MM-DD)", font=('', 8)).grid(row=0, column=2, sticky=tk.W, padx=5)

        # Jenis Transaksi
        ttk.Label(main_frame, text="Jenis Transaksi:").grid(row=1, column=0, sticky=tk.W, pady=5)
        self.jenis_var = tk.StringVar()
        self.jenis_combo = ttk.Combobox(main_frame, textvariable=self.jenis_var, state='readonly', width=30)
        self.jenis_combo['values'] = [f"{k} - {v}" for k, v in JENIS_TRANSAKSI.items()]
        self.jenis_combo.grid(row=1, column=1, columnspan=2, sticky=tk.W, pady=5)
        self.jenis_combo.bind('<<ComboboxSelected>>', self.on_jenis_change)

        # Uraian
        ttk.Label(main_frame, text="Uraian:").grid(row=2, column=0, sticky=tk.W, pady=5)
        self.uraian_var = tk.StringVar()
        self.uraian_entry = ttk.Entry(main_frame, textvariable=self.uraian_var, width=40)
        self.uraian_entry.grid(row=2, column=1, columnspan=2, sticky=tk.W, pady=5)

        # Jumlah
        ttk.Label(main_frame, text="Jumlah (Rp):").grid(row=3, column=0, sticky=tk.W, pady=5)
        self.jumlah_var = tk.StringVar()
        self.jumlah_entry = ttk.Entry(main_frame, textvariable=self.jumlah_var, width=20)
        self.jumlah_entry.grid(row=3, column=1, sticky=tk.W, pady=5)
        self.jumlah_entry.bind('<FocusOut>', self.format_jumlah)

        # Sumber/Tujuan Kas (untuk TRM, SDS, SDR)
        self.sumber_frame = ttk.LabelFrame(main_frame, text="Sumber/Tujuan Kas", padding=10)
        self.sumber_frame.grid(row=4, column=0, columnspan=3, sticky=tk.W+tk.E, pady=10)

        self.sumber_var = tk.StringVar(value='tunai')
        self.rb_tunai = ttk.Radiobutton(self.sumber_frame, text="Kas Tunai",
                                         variable=self.sumber_var, value='tunai')
        self.rb_tunai.pack(side=tk.LEFT, padx=10)
        self.rb_bank = ttk.Radiobutton(self.sumber_frame, text="Kas Bank",
                                        variable=self.sumber_var, value='bank')
        self.rb_bank.pack(side=tk.LEFT, padx=10)

        # Pos Anggaran
        self.pos_frame = ttk.LabelFrame(main_frame, text="Pos Anggaran", padding=10)
        self.pos_frame.grid(row=5, column=0, columnspan=3, sticky=tk.W+tk.E, pady=10)

        self.pos_var = tk.StringVar()
        self.pos_combo = ttk.Combobox(self.pos_frame, textvariable=self.pos_var, state='readonly', width=40)
        self.pos_combo.pack(fill=tk.X)
        self.load_pos_options()

        self.pos_label = ttk.Label(self.pos_frame, text="", font=('', 8))
        self.pos_label.pack(anchor=tk.W, pady=5)

        # Tombol
        btn_frame = ttk.Frame(main_frame)
        btn_frame.grid(row=6, column=0, columnspan=3, pady=20)

        ttk.Button(btn_frame, text="Simpan", command=self.save).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Batal", command=self.destroy).pack(side=tk.LEFT, padx=5)

        # Set initial visibility
        self.sumber_frame.grid_remove()
        self.pos_frame.grid_remove()

    def load_pos_options(self):
        """Memuat daftar pos anggaran"""
        pos_list = self.db.get_all_pos()
        self.pos_data = {f"{p['kode_pos']} - {p['nama_pos']}": p['id'] for p in pos_list}
        self.pos_combo['values'] = ['(Tidak ada pos)'] + list(self.pos_data.keys())
        self.pos_combo.set('(Tidak ada pos)')

    def on_jenis_change(self, event=None):
        """Handler ketika jenis transaksi berubah"""
        jenis = self.jenis_var.get()[:3] if self.jenis_var.get() else ''

        # Tampilkan/sembunyikan sumber kas
        if jenis in ('TRM', 'SDS', 'SDR'):
            self.sumber_frame.grid()
        else:
            self.sumber_frame.grid_remove()

        # Tampilkan/sembunyikan pos anggaran
        if jenis in ('TRM', 'BLJ', 'TRF', 'SDS', 'SDR'):
            self.pos_frame.grid()
            if jenis == 'TRM':
                self.pos_label.config(text="Opsional - pilih jika dana dialokasikan ke pos tertentu")
            else:
                self.pos_label.config(text="Wajib - pilih pos untuk transaksi ini")
        else:
            self.pos_frame.grid_remove()

    def format_jumlah(self, event=None):
        """Format jumlah dengan separator ribuan"""
        try:
            nilai = parse_rupiah(self.jumlah_var.get())
            if nilai > 0:
                self.jumlah_var.set(f"{nilai:,.0f}".replace(",", "."))
        except:
            pass

    def load_data(self):
        """Memuat data transaksi untuk edit"""
        t = self.transaksi
        self.tanggal_var.set(t['tanggal'])

        # Set jenis transaksi
        jenis = t['jenis_transaksi']
        for i, k in enumerate(JENIS_TRANSAKSI.keys()):
            if k == jenis:
                self.jenis_combo.current(i)
                break

        self.uraian_var.set(t['uraian'])
        self.jumlah_var.set(f"{t['jumlah']:,.0f}".replace(",", "."))

        if t['sumber_kas']:
            self.sumber_var.set(t['sumber_kas'])

        if t['pos_id']:
            pos = self.db.get_pos_by_id(t['pos_id'])
            if pos:
                self.pos_combo.set(f"{pos['kode_pos']} - {pos['nama_pos']}")

        self.on_jenis_change()

    def save(self):
        """Menyimpan transaksi"""
        # Validasi tanggal
        tanggal = self.tanggal_var.get().strip()
        try:
            datetime.strptime(tanggal, "%Y-%m-%d")
        except ValueError:
            messagebox.showerror("Error", "Format tanggal tidak valid. Gunakan YYYY-MM-DD")
            return

        # Validasi jenis
        jenis = self.jenis_var.get()[:3] if self.jenis_var.get() else ''
        if not jenis:
            messagebox.showerror("Error", "Pilih jenis transaksi")
            return

        # Validasi uraian
        uraian = self.uraian_var.get().strip()
        if not uraian:
            messagebox.showerror("Error", "Uraian tidak boleh kosong")
            return

        # Validasi jumlah
        jumlah = parse_rupiah(self.jumlah_var.get())
        if jumlah <= 0:
            messagebox.showerror("Error", "Jumlah harus lebih besar dari 0")
            return

        # Sumber kas
        sumber_kas = None
        if jenis in ('TRM', 'SDS', 'SDR'):
            sumber_kas = self.sumber_var.get()

        # Pos anggaran
        pos_id = None
        arah_pos = None
        pos_selected = self.pos_combo.get()

        if jenis in ('BLJ', 'TRF', 'SDS', 'SDR'):
            # Wajib pilih pos
            if pos_selected == '(Tidak ada pos)' or not pos_selected:
                messagebox.showerror("Error", f"Transaksi {JENIS_TRANSAKSI[jenis]} wajib memilih pos anggaran")
                return
            pos_id = self.pos_data.get(pos_selected)
            arah_pos = 'keluar'
        elif jenis == 'TRM' and pos_selected != '(Tidak ada pos)' and pos_selected:
            pos_id = self.pos_data.get(pos_selected)
            arah_pos = 'masuk'

        # Validasi saldo
        exclude_id = self.transaksi['id'] if self.transaksi else None
        valid, msg = self.db.validate_transaksi(jenis, jumlah, sumber_kas, exclude_id)
        if not valid:
            messagebox.showerror("Error", msg)
            return

        # Simpan
        self.result = {
            'tanggal': tanggal,
            'uraian': uraian,
            'jenis': jenis,
            'jumlah': jumlah,
            'sumber_kas': sumber_kas,
            'pos_id': pos_id,
            'arah_pos': arah_pos
        }
        self.destroy()


class PosDialog(tk.Toplevel):
    """Dialog untuk tambah/edit pos anggaran"""

    def __init__(self, parent, db, pos=None):
        super().__init__(parent)
        self.db = db
        self.pos = pos
        self.result = None

        self.title("Edit Pos Anggaran" if pos else "Tambah Pos Anggaran")
        self.geometry("450x350")
        self.resizable(False, False)
        self.transient(parent)
        self.grab_set()

        self.create_widgets()

        if pos:
            self.load_data()

        self.center_window()

    def center_window(self):
        """Posisikan window di tengah"""
        self.update_idletasks()
        x = (self.winfo_screenwidth() - self.winfo_width()) // 2
        y = (self.winfo_screenheight() - self.winfo_height()) // 2
        self.geometry(f"+{x}+{y}")

    def create_widgets(self):
        """Membuat widget dialog"""
        main_frame = ttk.Frame(self, padding=20)
        main_frame.pack(fill=tk.BOTH, expand=True)

        # Kode Pos
        ttk.Label(main_frame, text="Kode Pos:").grid(row=0, column=0, sticky=tk.W, pady=5)
        self.kode_var = tk.StringVar()
        self.kode_entry = ttk.Entry(main_frame, textvariable=self.kode_var, width=10)
        self.kode_entry.grid(row=0, column=1, sticky=tk.W, pady=5)
        ttk.Label(main_frame, text="(maks 5 karakter)", font=('', 8)).grid(row=0, column=2, sticky=tk.W, padx=5)

        # Nama Pos
        ttk.Label(main_frame, text="Nama Pos:").grid(row=1, column=0, sticky=tk.W, pady=5)
        self.nama_var = tk.StringVar()
        self.nama_entry = ttk.Entry(main_frame, textvariable=self.nama_var, width=35)
        self.nama_entry.grid(row=1, column=1, columnspan=2, sticky=tk.W, pady=5)

        # Target Anggaran
        ttk.Label(main_frame, text="Target Anggaran:").grid(row=2, column=0, sticky=tk.W, pady=5)
        self.target_var = tk.StringVar(value="0")
        self.target_entry = ttk.Entry(main_frame, textvariable=self.target_var, width=20)
        self.target_entry.grid(row=2, column=1, sticky=tk.W, pady=5)
        self.target_entry.bind('<FocusOut>', self.format_target)

        # Keterangan
        ttk.Label(main_frame, text="Keterangan:").grid(row=3, column=0, sticky=tk.NW, pady=5)
        self.keterangan_text = tk.Text(main_frame, width=35, height=5)
        self.keterangan_text.grid(row=3, column=1, columnspan=2, sticky=tk.W, pady=5)
        ttk.Label(main_frame, text="(catatan musyawarah, dll)", font=('', 8)).grid(row=4, column=1, sticky=tk.W)

        # Tombol
        btn_frame = ttk.Frame(main_frame)
        btn_frame.grid(row=5, column=0, columnspan=3, pady=20)

        ttk.Button(btn_frame, text="Simpan", command=self.save).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Batal", command=self.destroy).pack(side=tk.LEFT, padx=5)

    def format_target(self, event=None):
        """Format target dengan separator ribuan"""
        try:
            nilai = parse_rupiah(self.target_var.get())
            if nilai >= 0:
                self.target_var.set(f"{nilai:,.0f}".replace(",", "."))
        except:
            pass

    def load_data(self):
        """Memuat data pos untuk edit"""
        self.kode_var.set(self.pos['kode_pos'])
        self.nama_var.set(self.pos['nama_pos'])
        self.target_var.set(f"{self.pos['target_anggaran']:,.0f}".replace(",", "."))
        self.keterangan_text.insert('1.0', self.pos['keterangan'] or '')

    def save(self):
        """Menyimpan pos"""
        kode = self.kode_var.get().strip().upper()
        if not kode:
            messagebox.showerror("Error", "Kode pos tidak boleh kosong")
            return
        if len(kode) > 5:
            messagebox.showerror("Error", "Kode pos maksimal 5 karakter")
            return

        # Cek duplikat kode
        existing = self.db.fetchone('SELECT id FROM pos_anggaran WHERE kode_pos=?', (kode,))
        if existing:
            if not self.pos or existing['id'] != self.pos['id']:
                messagebox.showerror("Error", f"Kode pos '{kode}' sudah digunakan")
                return

        nama = self.nama_var.get().strip()
        if not nama:
            messagebox.showerror("Error", "Nama pos tidak boleh kosong")
            return

        target = parse_rupiah(self.target_var.get())
        keterangan = self.keterangan_text.get('1.0', tk.END).strip()

        self.result = {
            'kode_pos': kode,
            'nama_pos': nama,
            'target_anggaran': target,
            'keterangan': keterangan
        }
        self.destroy()


class BukuKasApp:
    """Aplikasi utama Buku Kas Kelompok"""

    def __init__(self, root):
        self.root = root
        self.db = Database(DB_FILE)

        self.setup_window()
        self.create_header()
        self.create_tabs()
        self.refresh_all()

    def setup_window(self):
        """Setup window utama"""
        self.root.title("BUKU KAS KELOMPOK")
        self.root.geometry(f"{WINDOW_WIDTH}x{WINDOW_HEIGHT}")

        # Style
        style = ttk.Style()
        style.theme_use('clam')

        # Configure Treeview
        style.configure("Treeview", rowheight=25, font=('', 10))
        style.configure("Treeview.Heading", font=('', 10, 'bold'))

        # Configure Labels
        style.configure("Header.TLabel", font=('', 14, 'bold'))
        style.configure("Saldo.TLabel", font=('', 12, 'bold'))
        style.configure("SaldoValue.TLabel", font=('', 14, 'bold'))

    def create_header(self):
        """Membuat header dengan saldo"""
        header_frame = ttk.Frame(self.root, padding=10)
        header_frame.pack(fill=tk.X)

        # Judul
        pengaturan = self.db.get_pengaturan()
        nama_kelompok = pengaturan['nama_kelompok'] if pengaturan else 'Nama Kelompok'

        self.title_label = ttk.Label(header_frame, text=f"BUKU KAS KELOMPOK - {nama_kelompok}",
                                      style="Header.TLabel")
        self.title_label.pack(side=tk.TOP, pady=5)

        # Frame saldo
        saldo_frame = ttk.Frame(header_frame)
        saldo_frame.pack(fill=tk.X, pady=10)

        # Saldo Tunai
        tunai_frame = ttk.LabelFrame(saldo_frame, text="SALDO TUNAI", padding=10)
        tunai_frame.pack(side=tk.LEFT, expand=True, fill=tk.X, padx=5)
        self.saldo_tunai_label = ttk.Label(tunai_frame, text="Rp 0", style="SaldoValue.TLabel",
                                            foreground='#28a745')
        self.saldo_tunai_label.pack()

        # Saldo Bank
        bank_frame = ttk.LabelFrame(saldo_frame, text="SALDO BANK", padding=10)
        bank_frame.pack(side=tk.LEFT, expand=True, fill=tk.X, padx=5)
        self.saldo_bank_label = ttk.Label(bank_frame, text="Rp 0", style="SaldoValue.TLabel",
                                           foreground='#007bff')
        self.saldo_bank_label.pack()

        # Total Kas
        total_frame = ttk.LabelFrame(saldo_frame, text="TOTAL KAS", padding=10)
        total_frame.pack(side=tk.LEFT, expand=True, fill=tk.X, padx=5)
        self.saldo_total_label = ttk.Label(total_frame, text="Rp 0", style="SaldoValue.TLabel",
                                            foreground='#6f42c1')
        self.saldo_total_label.pack()

    def create_tabs(self):
        """Membuat tab navigasi"""
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        # Tab Buku Kas
        self.tab_buku_kas = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(self.tab_buku_kas, text="Buku Kas")
        self.create_tab_buku_kas()

        # Tab Pos Anggaran
        self.tab_pos = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(self.tab_pos, text="Pos Anggaran")
        self.create_tab_pos()

        # Tab Rekonsiliasi
        self.tab_rekon = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(self.tab_rekon, text="Rekonsiliasi")
        self.create_tab_rekon()

        # Tab Laporan Bulanan
        self.tab_laporan = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(self.tab_laporan, text="Laporan Bulanan")
        self.create_tab_laporan()

        # Tab Pengaturan
        self.tab_pengaturan = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(self.tab_pengaturan, text="Pengaturan")
        self.create_tab_pengaturan()

    def create_tab_buku_kas(self):
        """Membuat konten tab Buku Kas"""
        # Toolbar
        toolbar = ttk.Frame(self.tab_buku_kas)
        toolbar.pack(fill=tk.X, pady=5)

        ttk.Button(toolbar, text="+ Tambah Transaksi", command=self.add_transaksi).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="Edit", command=self.edit_transaksi).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="Hapus", command=self.delete_transaksi).pack(side=tk.LEFT, padx=2)

        ttk.Separator(toolbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=10)

        # Filter tanggal
        ttk.Label(toolbar, text="Dari:").pack(side=tk.LEFT, padx=2)
        self.filter_dari_var = tk.StringVar()
        ttk.Entry(toolbar, textvariable=self.filter_dari_var, width=12).pack(side=tk.LEFT, padx=2)

        ttk.Label(toolbar, text="Sampai:").pack(side=tk.LEFT, padx=2)
        self.filter_sampai_var = tk.StringVar()
        ttk.Entry(toolbar, textvariable=self.filter_sampai_var, width=12).pack(side=tk.LEFT, padx=2)

        ttk.Button(toolbar, text="Filter", command=self.filter_transaksi).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="Reset", command=self.reset_filter).pack(side=tk.LEFT, padx=2)

        # Treeview transaksi
        columns = ('no', 'tanggal', 'uraian', 'jenis', 'pos', 'jumlah', 'saldo_tunai', 'saldo_bank')
        self.tree_transaksi = ttk.Treeview(self.tab_buku_kas, columns=columns, show='headings', height=15)

        self.tree_transaksi.heading('no', text='No')
        self.tree_transaksi.heading('tanggal', text='Tanggal')
        self.tree_transaksi.heading('uraian', text='Uraian')
        self.tree_transaksi.heading('jenis', text='Jenis')
        self.tree_transaksi.heading('pos', text='Pos')
        self.tree_transaksi.heading('jumlah', text='Jumlah')
        self.tree_transaksi.heading('saldo_tunai', text='Saldo Tunai')
        self.tree_transaksi.heading('saldo_bank', text='Saldo Bank')

        self.tree_transaksi.column('no', width=40, anchor=tk.CENTER)
        self.tree_transaksi.column('tanggal', width=90, anchor=tk.CENTER)
        self.tree_transaksi.column('uraian', width=250)
        self.tree_transaksi.column('jenis', width=60, anchor=tk.CENTER)
        self.tree_transaksi.column('pos', width=100)
        self.tree_transaksi.column('jumlah', width=120, anchor=tk.E)
        self.tree_transaksi.column('saldo_tunai', width=120, anchor=tk.E)
        self.tree_transaksi.column('saldo_bank', width=120, anchor=tk.E)

        # Scrollbar
        scrollbar = ttk.Scrollbar(self.tab_buku_kas, orient=tk.VERTICAL, command=self.tree_transaksi.yview)
        self.tree_transaksi.configure(yscrollcommand=scrollbar.set)

        self.tree_transaksi.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        # Tags untuk warna baris
        self.tree_transaksi.tag_configure('masuk', background=WARNA_TRANSAKSI['masuk'])
        self.tree_transaksi.tag_configure('keluar', background=WARNA_TRANSAKSI['keluar'])
        self.tree_transaksi.tag_configure('pindah', background=WARNA_TRANSAKSI['pindah'])
        self.tree_transaksi.tag_configure('setor', background=WARNA_TRANSAKSI['setor'])

    def create_tab_pos(self):
        """Membuat konten tab Pos Anggaran"""
        # Toolbar
        toolbar = ttk.Frame(self.tab_pos)
        toolbar.pack(fill=tk.X, pady=5)

        ttk.Button(toolbar, text="+ Tambah Pos", command=self.add_pos).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="Edit", command=self.edit_pos).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="Hapus", command=self.delete_pos).pack(side=tk.LEFT, padx=2)

        # Treeview pos
        columns = ('no', 'kode', 'nama', 'target', 'masuk', 'keluar', 'saldo', 'persen')
        self.tree_pos = ttk.Treeview(self.tab_pos, columns=columns, show='headings', height=15)

        self.tree_pos.heading('no', text='No')
        self.tree_pos.heading('kode', text='Kode')
        self.tree_pos.heading('nama', text='Nama Pos')
        self.tree_pos.heading('target', text='Target')
        self.tree_pos.heading('masuk', text='Masuk')
        self.tree_pos.heading('keluar', text='Keluar')
        self.tree_pos.heading('saldo', text='Saldo Pos')
        self.tree_pos.heading('persen', text='% Terpakai')

        self.tree_pos.column('no', width=40, anchor=tk.CENTER)
        self.tree_pos.column('kode', width=60, anchor=tk.CENTER)
        self.tree_pos.column('nama', width=200)
        self.tree_pos.column('target', width=120, anchor=tk.E)
        self.tree_pos.column('masuk', width=120, anchor=tk.E)
        self.tree_pos.column('keluar', width=120, anchor=tk.E)
        self.tree_pos.column('saldo', width=120, anchor=tk.E)
        self.tree_pos.column('persen', width=80, anchor=tk.CENTER)

        # Scrollbar
        scrollbar = ttk.Scrollbar(self.tab_pos, orient=tk.VERTICAL, command=self.tree_pos.yview)
        self.tree_pos.configure(yscrollcommand=scrollbar.set)

        self.tree_pos.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        # Tag untuk over budget
        self.tree_pos.tag_configure('over', background='#f8d7da')

    def create_tab_rekon(self):
        """Membuat konten tab Rekonsiliasi"""
        # Frame input
        input_frame = ttk.LabelFrame(self.tab_rekon, text="Input Rekonsiliasi", padding=15)
        input_frame.pack(fill=tk.X, pady=10)

        # Grid layout
        ttk.Label(input_frame, text="Tanggal:").grid(row=0, column=0, sticky=tk.W, pady=5)
        self.rekon_tanggal_var = tk.StringVar(value=date.today().strftime("%Y-%m-%d"))
        ttk.Entry(input_frame, textvariable=self.rekon_tanggal_var, width=15).grid(row=0, column=1, sticky=tk.W, pady=5)

        # Perbandingan
        compare_frame = ttk.Frame(input_frame)
        compare_frame.grid(row=1, column=0, columnspan=4, pady=10)

        # Kondisi Fisik
        fisik_frame = ttk.LabelFrame(compare_frame, text="KONDISI FISIK", padding=10)
        fisik_frame.pack(side=tk.LEFT, padx=10)

        ttk.Label(fisik_frame, text="Tunai Fisik:").grid(row=0, column=0, sticky=tk.W, pady=3)
        self.rekon_tunai_fisik_var = tk.StringVar(value="0")
        ttk.Entry(fisik_frame, textvariable=self.rekon_tunai_fisik_var, width=15).grid(row=0, column=1, pady=3)

        ttk.Label(fisik_frame, text="Bank Fisik:").grid(row=1, column=0, sticky=tk.W, pady=3)
        self.rekon_bank_fisik_var = tk.StringVar(value="0")
        ttk.Entry(fisik_frame, textvariable=self.rekon_bank_fisik_var, width=15).grid(row=1, column=1, pady=3)

        # Kondisi Sistem
        sistem_frame = ttk.LabelFrame(compare_frame, text="KONDISI SISTEM", padding=10)
        sistem_frame.pack(side=tk.LEFT, padx=10)

        ttk.Label(sistem_frame, text="Tunai Sistem:").grid(row=0, column=0, sticky=tk.W, pady=3)
        self.rekon_tunai_sistem_label = ttk.Label(sistem_frame, text="Rp 0")
        self.rekon_tunai_sistem_label.grid(row=0, column=1, pady=3)

        ttk.Label(sistem_frame, text="Bank Sistem:").grid(row=1, column=0, sticky=tk.W, pady=3)
        self.rekon_bank_sistem_label = ttk.Label(sistem_frame, text="Rp 0")
        self.rekon_bank_sistem_label.grid(row=1, column=1, pady=3)

        # Selisih
        selisih_frame = ttk.LabelFrame(compare_frame, text="SELISIH", padding=10)
        selisih_frame.pack(side=tk.LEFT, padx=10)

        ttk.Label(selisih_frame, text="Selisih Tunai:").grid(row=0, column=0, sticky=tk.W, pady=3)
        self.rekon_selisih_tunai_label = ttk.Label(selisih_frame, text="Rp 0")
        self.rekon_selisih_tunai_label.grid(row=0, column=1, pady=3)

        ttk.Label(selisih_frame, text="Selisih Bank:").grid(row=1, column=0, sticky=tk.W, pady=3)
        self.rekon_selisih_bank_label = ttk.Label(selisih_frame, text="Rp 0")
        self.rekon_selisih_bank_label.grid(row=1, column=1, pady=3)

        # Keterangan
        ttk.Label(input_frame, text="Keterangan:").grid(row=2, column=0, sticky=tk.NW, pady=5)
        self.rekon_keterangan_text = tk.Text(input_frame, width=50, height=3)
        self.rekon_keterangan_text.grid(row=2, column=1, columnspan=3, sticky=tk.W, pady=5)

        # Tombol
        btn_frame = ttk.Frame(input_frame)
        btn_frame.grid(row=3, column=0, columnspan=4, pady=10)
        ttk.Button(btn_frame, text="Hitung Selisih", command=self.hitung_selisih).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Simpan Rekap", command=self.simpan_rekon).pack(side=tk.LEFT, padx=5)

        # Treeview riwayat
        ttk.Label(self.tab_rekon, text="Riwayat Rekonsiliasi:", font=('', 10, 'bold')).pack(anchor=tk.W, pady=5)

        columns = ('tanggal', 'tunai_fisik', 'tunai_sistem', 'selisih_tunai',
                   'bank_fisik', 'bank_sistem', 'selisih_bank', 'keterangan')
        self.tree_rekon = ttk.Treeview(self.tab_rekon, columns=columns, show='headings', height=8)

        self.tree_rekon.heading('tanggal', text='Tanggal')
        self.tree_rekon.heading('tunai_fisik', text='Tunai Fisik')
        self.tree_rekon.heading('tunai_sistem', text='Tunai Sistem')
        self.tree_rekon.heading('selisih_tunai', text='Selisih')
        self.tree_rekon.heading('bank_fisik', text='Bank Fisik')
        self.tree_rekon.heading('bank_sistem', text='Bank Sistem')
        self.tree_rekon.heading('selisih_bank', text='Selisih')
        self.tree_rekon.heading('keterangan', text='Keterangan')

        self.tree_rekon.column('tanggal', width=90, anchor=tk.CENTER)
        self.tree_rekon.column('tunai_fisik', width=100, anchor=tk.E)
        self.tree_rekon.column('tunai_sistem', width=100, anchor=tk.E)
        self.tree_rekon.column('selisih_tunai', width=80, anchor=tk.E)
        self.tree_rekon.column('bank_fisik', width=100, anchor=tk.E)
        self.tree_rekon.column('bank_sistem', width=100, anchor=tk.E)
        self.tree_rekon.column('selisih_bank', width=80, anchor=tk.E)
        self.tree_rekon.column('keterangan', width=150)

        self.tree_rekon.pack(fill=tk.BOTH, expand=True)

        # Tag untuk selisih
        self.tree_rekon.tag_configure('selisih', background='#fff3cd')

    def create_tab_laporan(self):
        """Membuat konten tab Laporan Bulanan"""
        # Filter
        filter_frame = ttk.Frame(self.tab_laporan)
        filter_frame.pack(fill=tk.X, pady=10)

        ttk.Label(filter_frame, text="Bulan:").pack(side=tk.LEFT, padx=5)
        self.laporan_bulan_var = tk.StringVar()
        bulan_combo = ttk.Combobox(filter_frame, textvariable=self.laporan_bulan_var, state='readonly', width=12)
        bulan_combo['values'] = ['Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni',
                                  'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember']
        bulan_combo.current(datetime.now().month - 1)
        bulan_combo.pack(side=tk.LEFT, padx=5)

        ttk.Label(filter_frame, text="Tahun:").pack(side=tk.LEFT, padx=5)
        self.laporan_tahun_var = tk.StringVar(value=str(datetime.now().year))
        tahun_combo = ttk.Combobox(filter_frame, textvariable=self.laporan_tahun_var, state='readonly', width=8)
        tahun_combo['values'] = [str(y) for y in range(2020, datetime.now().year + 2)]
        tahun_combo.pack(side=tk.LEFT, padx=5)

        ttk.Button(filter_frame, text="Tampilkan", command=self.tampilkan_laporan).pack(side=tk.LEFT, padx=10)
        ttk.Button(filter_frame, text="Export Excel", command=self.export_excel).pack(side=tk.LEFT, padx=5)

        # Frame laporan dengan scrollbar
        self.laporan_canvas = tk.Canvas(self.tab_laporan)
        scrollbar = ttk.Scrollbar(self.tab_laporan, orient=tk.VERTICAL, command=self.laporan_canvas.yview)
        self.laporan_frame = ttk.Frame(self.laporan_canvas)

        self.laporan_frame.bind("<Configure>", lambda e: self.laporan_canvas.configure(
            scrollregion=self.laporan_canvas.bbox("all")))

        self.laporan_canvas.create_window((0, 0), window=self.laporan_frame, anchor=tk.NW)
        self.laporan_canvas.configure(yscrollcommand=scrollbar.set)

        self.laporan_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    def create_tab_pengaturan(self):
        """Membuat konten tab Pengaturan"""
        form_frame = ttk.LabelFrame(self.tab_pengaturan, text="Pengaturan Kelompok", padding=20)
        form_frame.pack(fill=tk.X, pady=20, padx=50)

        # Nama Kelompok
        ttk.Label(form_frame, text="Nama Kelompok:").grid(row=0, column=0, sticky=tk.W, pady=10)
        self.pengaturan_nama_var = tk.StringVar()
        ttk.Entry(form_frame, textvariable=self.pengaturan_nama_var, width=40).grid(row=0, column=1, pady=10)

        # Alamat
        ttk.Label(form_frame, text="Alamat/Desa:").grid(row=1, column=0, sticky=tk.W, pady=10)
        self.pengaturan_alamat_var = tk.StringVar()
        ttk.Entry(form_frame, textvariable=self.pengaturan_alamat_var, width=40).grid(row=1, column=1, pady=10)

        # Nama Bendahara
        ttk.Label(form_frame, text="Nama Bendahara:").grid(row=2, column=0, sticky=tk.W, pady=10)
        self.pengaturan_bendahara_var = tk.StringVar()
        ttk.Entry(form_frame, textvariable=self.pengaturan_bendahara_var, width=40).grid(row=2, column=1, pady=10)

        # Tombol Simpan
        ttk.Button(form_frame, text="Simpan Pengaturan", command=self.simpan_pengaturan).grid(
            row=3, column=0, columnspan=2, pady=20)

        # Load data pengaturan
        self.load_pengaturan()

    def load_pengaturan(self):
        """Memuat data pengaturan"""
        pengaturan = self.db.get_pengaturan()
        if pengaturan:
            self.pengaturan_nama_var.set(pengaturan['nama_kelompok'] or '')
            self.pengaturan_alamat_var.set(pengaturan['alamat'] or '')
            self.pengaturan_bendahara_var.set(pengaturan['nama_bendahara'] or '')

    def simpan_pengaturan(self):
        """Menyimpan pengaturan"""
        nama = self.pengaturan_nama_var.get().strip()
        alamat = self.pengaturan_alamat_var.get().strip()
        bendahara = self.pengaturan_bendahara_var.get().strip()

        self.db.update_pengaturan(nama, alamat, bendahara)

        # Update judul
        self.title_label.config(text=f"BUKU KAS KELOMPOK - {nama}")

        messagebox.showinfo("Sukses", "Pengaturan berhasil disimpan")

    def refresh_saldo(self):
        """Refresh tampilan saldo"""
        tunai = self.db.get_saldo_tunai()
        bank = self.db.get_saldo_bank()
        total = tunai + bank

        self.saldo_tunai_label.config(text=format_rupiah(tunai))
        self.saldo_bank_label.config(text=format_rupiah(bank))
        self.saldo_total_label.config(text=format_rupiah(total))

        # Update rekonsiliasi
        self.rekon_tunai_sistem_label.config(text=format_rupiah(tunai))
        self.rekon_bank_sistem_label.config(text=format_rupiah(bank))

    def refresh_transaksi(self, tgl_dari=None, tgl_sampai=None):
        """Refresh tabel transaksi"""
        # Clear existing items
        for item in self.tree_transaksi.get_children():
            self.tree_transaksi.delete(item)

        transaksi = self.db.get_all_transaksi(tgl_dari, tgl_sampai)

        for i, t in enumerate(transaksi, 1):
            jenis = t['jenis_transaksi']

            # Tentukan tag warna
            if jenis in ('TRM', 'BNG'):
                tag = 'masuk'
            elif jenis in ('BLJ', 'TRF', 'ADM'):
                tag = 'keluar'
            elif jenis in ('STR', 'TRK'):
                tag = 'pindah'
            else:  # SDS, SDR
                tag = 'setor'

            pos_str = f"{t['kode_pos']}" if t['kode_pos'] else "-"

            self.tree_transaksi.insert('', tk.END, iid=str(t['id']), values=(
                i,
                t['tanggal'],
                t['uraian'],
                jenis,
                pos_str,
                format_rupiah(t['jumlah']),
                format_rupiah(t['saldo_tunai_setelah']),
                format_rupiah(t['saldo_bank_setelah'])
            ), tags=(tag,))

    def refresh_pos(self):
        """Refresh tabel pos anggaran"""
        # Clear existing items
        for item in self.tree_pos.get_children():
            self.tree_pos.delete(item)

        pos_list = self.db.get_all_pos()

        for i, p in enumerate(pos_list, 1):
            summary = self.db.get_pos_summary(p['id'])
            saldo = summary['masuk'] - summary['keluar']
            target = p['target_anggaran']
            persen = (summary['keluar'] / target * 100) if target > 0 else 0

            # Tag untuk over budget
            tag = 'over' if summary['keluar'] > target and target > 0 else ''

            self.tree_pos.insert('', tk.END, iid=str(p['id']), values=(
                i,
                p['kode_pos'],
                p['nama_pos'],
                format_rupiah(target),
                format_rupiah(summary['masuk']),
                format_rupiah(summary['keluar']),
                format_rupiah(saldo),
                f"{persen:.1f}%"
            ), tags=(tag,) if tag else ())

    def refresh_rekon(self):
        """Refresh tabel rekonsiliasi"""
        # Clear existing items
        for item in self.tree_rekon.get_children():
            self.tree_rekon.delete(item)

        rekon_list = self.db.get_all_rekonsiliasi()

        for r in rekon_list:
            # Tag jika ada selisih
            has_selisih = r['selisih_tunai'] != 0 or r['selisih_bank'] != 0
            tag = 'selisih' if has_selisih else ''

            self.tree_rekon.insert('', tk.END, values=(
                r['tanggal'],
                format_rupiah(r['tunai_fisik']),
                format_rupiah(r['tunai_sistem']),
                format_rupiah(r['selisih_tunai']),
                format_rupiah(r['bank_fisik']),
                format_rupiah(r['bank_sistem']),
                format_rupiah(r['selisih_bank']),
                r['keterangan'][:30] + '...' if len(r['keterangan']) > 30 else r['keterangan']
            ), tags=(tag,) if tag else ())

    def refresh_all(self):
        """Refresh semua tampilan"""
        self.refresh_saldo()
        self.refresh_transaksi()
        self.refresh_pos()
        self.refresh_rekon()

    def add_transaksi(self):
        """Tambah transaksi baru"""
        dialog = TransaksiDialog(self.root, self.db)
        self.root.wait_window(dialog)

        if dialog.result:
            r = dialog.result
            self.db.add_transaksi(r['tanggal'], r['uraian'], r['jenis'], r['jumlah'],
                                   r['sumber_kas'], r['pos_id'], r['arah_pos'])
            self.refresh_all()

    def edit_transaksi(self):
        """Edit transaksi yang dipilih"""
        selected = self.tree_transaksi.selection()
        if not selected:
            messagebox.showwarning("Peringatan", "Pilih transaksi yang akan diedit")
            return

        trans_id = int(selected[0])
        transaksi = self.db.get_transaksi_by_id(trans_id)

        dialog = TransaksiDialog(self.root, self.db, transaksi)
        self.root.wait_window(dialog)

        if dialog.result:
            r = dialog.result
            self.db.update_transaksi(trans_id, r['tanggal'], r['uraian'], r['jenis'], r['jumlah'],
                                      r['sumber_kas'], r['pos_id'], r['arah_pos'])
            self.refresh_all()

    def delete_transaksi(self):
        """Hapus transaksi yang dipilih"""
        selected = self.tree_transaksi.selection()
        if not selected:
            messagebox.showwarning("Peringatan", "Pilih transaksi yang akan dihapus")
            return

        if messagebox.askyesno("Konfirmasi", "Yakin ingin menghapus transaksi ini?"):
            trans_id = int(selected[0])
            self.db.delete_transaksi(trans_id)
            self.refresh_all()

    def filter_transaksi(self):
        """Filter transaksi berdasarkan tanggal"""
        tgl_dari = self.filter_dari_var.get().strip() or None
        tgl_sampai = self.filter_sampai_var.get().strip() or None
        self.refresh_transaksi(tgl_dari, tgl_sampai)

    def reset_filter(self):
        """Reset filter tanggal"""
        self.filter_dari_var.set('')
        self.filter_sampai_var.set('')
        self.refresh_transaksi()

    def add_pos(self):
        """Tambah pos anggaran baru"""
        dialog = PosDialog(self.root, self.db)
        self.root.wait_window(dialog)

        if dialog.result:
            r = dialog.result
            self.db.add_pos(r['kode_pos'], r['nama_pos'], r['target_anggaran'], r['keterangan'])
            self.refresh_pos()

    def edit_pos(self):
        """Edit pos yang dipilih"""
        selected = self.tree_pos.selection()
        if not selected:
            messagebox.showwarning("Peringatan", "Pilih pos yang akan diedit")
            return

        pos_id = int(selected[0])
        pos = self.db.get_pos_by_id(pos_id)

        dialog = PosDialog(self.root, self.db, pos)
        self.root.wait_window(dialog)

        if dialog.result:
            r = dialog.result
            self.db.update_pos(pos_id, r['kode_pos'], r['nama_pos'], r['target_anggaran'], r['keterangan'])
            self.refresh_pos()

    def delete_pos(self):
        """Hapus pos yang dipilih"""
        selected = self.tree_pos.selection()
        if not selected:
            messagebox.showwarning("Peringatan", "Pilih pos yang akan dihapus")
            return

        pos_id = int(selected[0])

        # Cek apakah pos memiliki transaksi
        if self.db.pos_has_transactions(pos_id):
            messagebox.showerror("Error", "Tidak dapat menghapus pos yang sudah memiliki transaksi")
            return

        if messagebox.askyesno("Konfirmasi", "Yakin ingin menghapus pos ini?"):
            self.db.delete_pos(pos_id)
            self.refresh_pos()

    def hitung_selisih(self):
        """Hitung selisih rekonsiliasi"""
        tunai_fisik = parse_rupiah(self.rekon_tunai_fisik_var.get())
        bank_fisik = parse_rupiah(self.rekon_bank_fisik_var.get())

        tunai_sistem = self.db.get_saldo_tunai()
        bank_sistem = self.db.get_saldo_bank()

        selisih_tunai = tunai_fisik - tunai_sistem
        selisih_bank = bank_fisik - bank_sistem

        self.rekon_selisih_tunai_label.config(text=format_rupiah(selisih_tunai),
                                               foreground='red' if selisih_tunai != 0 else 'green')
        self.rekon_selisih_bank_label.config(text=format_rupiah(selisih_bank),
                                              foreground='red' if selisih_bank != 0 else 'green')

    def simpan_rekon(self):
        """Simpan data rekonsiliasi"""
        tanggal = self.rekon_tanggal_var.get().strip()
        try:
            datetime.strptime(tanggal, "%Y-%m-%d")
        except ValueError:
            messagebox.showerror("Error", "Format tanggal tidak valid")
            return

        tunai_fisik = parse_rupiah(self.rekon_tunai_fisik_var.get())
        bank_fisik = parse_rupiah(self.rekon_bank_fisik_var.get())
        keterangan = self.rekon_keterangan_text.get('1.0', tk.END).strip()

        self.db.add_rekonsiliasi(tanggal, tunai_fisik, bank_fisik, keterangan)

        # Reset form
        self.rekon_tunai_fisik_var.set("0")
        self.rekon_bank_fisik_var.set("0")
        self.rekon_keterangan_text.delete('1.0', tk.END)
        self.rekon_selisih_tunai_label.config(text="Rp 0", foreground='black')
        self.rekon_selisih_bank_label.config(text="Rp 0", foreground='black')

        self.refresh_rekon()
        messagebox.showinfo("Sukses", "Rekonsiliasi berhasil disimpan")

    def tampilkan_laporan(self):
        """Tampilkan laporan bulanan"""
        # Clear frame laporan
        for widget in self.laporan_frame.winfo_children():
            widget.destroy()

        bulan_idx = ['Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni',
                     'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember']
        bulan = bulan_idx.index(self.laporan_bulan_var.get()) + 1
        tahun = int(self.laporan_tahun_var.get())

        data = self.db.get_laporan_bulanan(bulan, tahun)

        # Judul
        ttk.Label(self.laporan_frame, text=f"LAPORAN KEUANGAN BULAN {self.laporan_bulan_var.get().upper()} {tahun}",
                  font=('', 14, 'bold')).pack(pady=10)

        # A. Ringkasan Bulan
        ringkasan_frame = ttk.LabelFrame(self.laporan_frame, text="A. RINGKASAN BULAN", padding=10)
        ringkasan_frame.pack(fill=tk.X, padx=20, pady=10)

        info = [
            ("Saldo Awal Tunai", data['saldo_awal_tunai']),
            ("Saldo Awal Bank", data['saldo_awal_bank']),
            ("Total Penerimaan (TRM + BNG)", data['total_penerimaan']),
            ("Total Pengeluaran (BLJ + TRF + ADM)", data['total_pengeluaran']),
            ("Total Setoran Desa/Daerah", data['total_setoran']),
            ("Saldo Akhir Tunai", data['saldo_akhir_tunai']),
            ("Saldo Akhir Bank", data['saldo_akhir_bank']),
        ]

        for i, (label, nilai) in enumerate(info):
            ttk.Label(ringkasan_frame, text=f"{label}:").grid(row=i, column=0, sticky=tk.W, pady=2)
            ttk.Label(ringkasan_frame, text=format_rupiah(nilai), font=('', 10, 'bold')).grid(
                row=i, column=1, sticky=tk.E, pady=2, padx=20)

        # B. Rekapitulasi Per Pos
        pos_frame = ttk.LabelFrame(self.laporan_frame, text="B. REKAPITULASI PER POS", padding=10)
        pos_frame.pack(fill=tk.X, padx=20, pady=10)

        columns = ('kode', 'nama', 'target', 'masuk', 'keluar', 'saldo')
        tree_pos = ttk.Treeview(pos_frame, columns=columns, show='headings', height=6)

        tree_pos.heading('kode', text='Kode')
        tree_pos.heading('nama', text='Nama Pos')
        tree_pos.heading('target', text='Target')
        tree_pos.heading('masuk', text='Masuk Bulan Ini')
        tree_pos.heading('keluar', text='Keluar Bulan Ini')
        tree_pos.heading('saldo', text='Saldo Pos')

        tree_pos.column('kode', width=60)
        tree_pos.column('nama', width=200)
        tree_pos.column('target', width=100, anchor=tk.E)
        tree_pos.column('masuk', width=120, anchor=tk.E)
        tree_pos.column('keluar', width=120, anchor=tk.E)
        tree_pos.column('saldo', width=100, anchor=tk.E)

        for p in data['rekap_pos']:
            saldo = p['masuk'] - p['keluar']
            tree_pos.insert('', tk.END, values=(
                p['kode_pos'],
                p['nama_pos'],
                format_rupiah(p['target_anggaran']),
                format_rupiah(p['masuk']),
                format_rupiah(p['keluar']),
                format_rupiah(saldo)
            ))

        tree_pos.pack(fill=tk.X)

        # C. Daftar Setoran
        if data['daftar_setoran']:
            setor_frame = ttk.LabelFrame(self.laporan_frame, text="C. DAFTAR SETORAN KE DESA/DAERAH", padding=10)
            setor_frame.pack(fill=tk.X, padx=20, pady=10)

            columns = ('tanggal', 'uraian', 'tujuan', 'pos', 'jumlah')
            tree_setor = ttk.Treeview(setor_frame, columns=columns, show='headings', height=5)

            tree_setor.heading('tanggal', text='Tanggal')
            tree_setor.heading('uraian', text='Uraian')
            tree_setor.heading('tujuan', text='Tujuan')
            tree_setor.heading('pos', text='Pos')
            tree_setor.heading('jumlah', text='Jumlah')

            tree_setor.column('tanggal', width=90)
            tree_setor.column('uraian', width=250)
            tree_setor.column('tujuan', width=80)
            tree_setor.column('pos', width=100)
            tree_setor.column('jumlah', width=120, anchor=tk.E)

            for s in data['daftar_setoran']:
                tujuan = "Desa" if s['jenis_transaksi'] == 'SDS' else "Daerah"
                pos_str = s['kode_pos'] if s['kode_pos'] else "-"
                tree_setor.insert('', tk.END, values=(
                    s['tanggal'],
                    s['uraian'],
                    tujuan,
                    pos_str,
                    format_rupiah(s['jumlah'])
                ))

            tree_setor.pack(fill=tk.X)

        # Store data for export
        self.laporan_data = data
        self.laporan_bulan = bulan
        self.laporan_tahun = tahun

    def export_excel(self):
        """Export laporan ke Excel"""
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
        except ImportError:
            messagebox.showerror("Error", "Module openpyxl tidak ditemukan.\nInstall dengan: pip install openpyxl")
            return

        if not hasattr(self, 'laporan_data'):
            messagebox.showwarning("Peringatan", "Tampilkan laporan terlebih dahulu")
            return

        # Dialog save file
        filename = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfilename=f"Laporan_Kas_{self.laporan_bulan:02d}_{self.laporan_tahun}.xlsx"
        )

        if not filename:
            return

        # Buat workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Laporan Bulanan"

        # Style
        header_font = Font(bold=True, size=14)
        subheader_font = Font(bold=True, size=11)
        bold_font = Font(bold=True)
        currency_align = Alignment(horizontal='right')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        bulan_nama = ['Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni',
                      'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember']

        # Judul
        pengaturan = self.db.get_pengaturan()
        ws['A1'] = f"LAPORAN KEUANGAN - {pengaturan['nama_kelompok']}"
        ws['A1'].font = header_font
        ws['A2'] = f"Bulan {bulan_nama[self.laporan_bulan - 1]} {self.laporan_tahun}"
        ws['A2'].font = subheader_font

        row = 4

        # A. Ringkasan
        ws[f'A{row}'] = "A. RINGKASAN BULAN"
        ws[f'A{row}'].font = subheader_font
        row += 1

        data = self.laporan_data
        ringkasan = [
            ("Saldo Awal Tunai", data['saldo_awal_tunai']),
            ("Saldo Awal Bank", data['saldo_awal_bank']),
            ("Total Penerimaan", data['total_penerimaan']),
            ("Total Pengeluaran", data['total_pengeluaran']),
            ("Total Setoran Desa/Daerah", data['total_setoran']),
            ("Saldo Akhir Tunai", data['saldo_akhir_tunai']),
            ("Saldo Akhir Bank", data['saldo_akhir_bank']),
        ]

        for label, nilai in ringkasan:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = nilai
            ws[f'B{row}'].number_format = '#,##0'
            ws[f'B{row}'].alignment = currency_align
            row += 1

        row += 2

        # B. Rekapitulasi Per Pos
        ws[f'A{row}'] = "B. REKAPITULASI PER POS"
        ws[f'A{row}'].font = subheader_font
        row += 1

        headers = ['Kode', 'Nama Pos', 'Target', 'Masuk', 'Keluar', 'Saldo']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = bold_font
            cell.border = thin_border
        row += 1

        for p in data['rekap_pos']:
            saldo = p['masuk'] - p['keluar']
            values = [p['kode_pos'], p['nama_pos'], p['target_anggaran'], p['masuk'], p['keluar'], saldo]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.border = thin_border
                if col >= 3:
                    cell.number_format = '#,##0'
                    cell.alignment = currency_align
            row += 1

        row += 2

        # C. Daftar Setoran
        if data['daftar_setoran']:
            ws[f'A{row}'] = "C. DAFTAR SETORAN KE DESA/DAERAH"
            ws[f'A{row}'].font = subheader_font
            row += 1

            headers = ['Tanggal', 'Uraian', 'Tujuan', 'Pos', 'Jumlah']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = bold_font
                cell.border = thin_border
            row += 1

            for s in data['daftar_setoran']:
                tujuan = "Desa" if s['jenis_transaksi'] == 'SDS' else "Daerah"
                pos_str = s['kode_pos'] if s['kode_pos'] else "-"
                values = [s['tanggal'], s['uraian'], tujuan, pos_str, s['jumlah']]
                for col, val in enumerate(values, 1):
                    cell = ws.cell(row=row, column=col, value=val)
                    cell.border = thin_border
                    if col == 5:
                        cell.number_format = '#,##0'
                        cell.alignment = currency_align
                row += 1

        # Adjust column widths
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 15

        # Save
        wb.save(filename)
        messagebox.showinfo("Sukses", f"Laporan berhasil diekspor ke:\n{filename}")


def main():
    """Fungsi utama"""
    root = tk.Tk()
    app = BukuKasApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
